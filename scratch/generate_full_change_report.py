# -*- coding: utf-8 -*-
"""
Script sinh file Word: Bao_cao_chi_tiet_cac_file_them_sua_va_tinh_nang.docx
Báo cáo kỹ thuật chi tiết toàn bộ các file đã thêm, sửa, mục đích từng file, chi tiết sửa đổi code.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, col_widths, header_bg="1E3A8A"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for idx, cell in enumerate(row.cells):
                set_cell_background(cell, header_bg)
                set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
        else:
            bg_color = "F8FAFC" if i % 2 == 1 else "FFFFFF"
            for idx, cell in enumerate(row.cells):
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(30, 41, 59)

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

def add_callout(doc, title, text, bg_hex="EFF6FF", border_hex="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    cell.width = Inches(6.8)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(title + "\n")
    run_t.font.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(30, 58, 138)
    
    run_b = p.add_run(text)
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    cell.width = Inches(6.8)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    for r in h.runs:
        r.font.name = "Arial"
        if level == 1:
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)
        elif level == 2:
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(37, 99, 235)
        elif level == 3:
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(71, 85, 105)
    return h

def main():
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BÁO CÁO KỸ THUẬT CHI TIẾT")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Liệt Kê Toàn Bộ Các File Đã Thêm, Sửa, Mục Đích & Nội Dung Thay Đổi Chi Tiết\nHệ Thống Gia Phả Đa Dòng Tộc (Multi-Family Genealogy)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    add_callout(
        doc,
        "MỤC ĐÍCH TÀI LIỆU",
        "Tài liệu này tổng hợp toàn bộ các thay đổi trong mã nguồn (Backend & Frontend) từ đầu đến thời điểm hiện tại: "
        "từ cơ chế một tài khoản thuộc đồng thời 2 dòng họ (Nội/Ngoại), chọn dòng tộc độc lập, rời dòng họ an toàn, "
        "hiển thị trạng thái chưa nối phả hệ trong cây, dọn dẹp các emoji/icon thừa, đến tính năng mới nhất: "
        "phân loại chính xác Dòng họ Nội & Dòng họ Ngoại, truy vấn thông tin phả hệ và hiển thị danh sách thành viên chi tiết.",
        bg_hex="F0FDF4",
        border_hex="22C55E"
    )
    
    # PHẦN 1
    add_h(doc, "1. TỔNG QUAN CÁC HẠNG MỤC ĐÃ THỰC HIỆN", level=1)
    
    summary_items = [
        ("Tính năng Đa Dòng Tộc (Multi-Clan Membership):", 
         "Cho phép 1 tài khoản tham gia đồng thời cả 2 dòng tộc (Bên Nội và Bên Ngoại), độc lập vai trò (Family Admin ở họ A, Member ở họ B), không yêu cầu đăng xuất để đổi tài khoản."),
        ("Quy trình Gia nhập Dòng họ mới (Join Clan Workflow):",
         "Cung cấp giao diện gửi yêu cầu (ClanJoinRequest) có khai báo quan hệ, ghi chú xác minh; quy trình phê duyệt hồ sơ cho Family Admin kèm liên kết hồ sơ Person."),
        ("Tính năng Rời Dòng Họ (Leave Clan):",
         "Cho phép người dùng chủ động rời khỏi dòng họ khi vào nhầm; modal xác nhận bảo mật và cập nhật cơ sở dữ liệu an toàn."),
        ("Hiển thị Thành Viên Chưa Nối Phả Hệ Trên Cây Gia Phả:",
         "Khi gia nhập thành công nhưng chưa xác nhận quan hệ, thành viên xuất hiện trên cây gia phả dưới dạng node chưa nối dây; khi xác nhận quan hệ, hệ thống tự động vẽ đường nối."),
        ("Dọn Dẹp & Chuẩn Hóa Giao Diện (Xoá icon thừa):",
         "Xoá các emoji không cần thiết (🏛️, 🔄, 🌳, 📩, ℹ️, 📋) trên trang 'Dòng tộc của tôi' và các modal quản trị; sửa lỗi tràn và chèn lấn giữa ClanSelector và thanh tìm kiếm TopBar."),
        ("Mạng Lưới Liên Họ - Phân Loại & Hiển Thị Thành Viên Dòng Họ Nội / Ngoại:",
         "Cập nhật API GET /networks/families: phân loại dòng tộc chính xác theo quan hệ lineage_type; truy vấn đầy đủ thông tin Thủy tổ, Quê quán, Lịch sử và danh sách thành viên thực tế (Person: họ tên, đời, giới tính, ngày sinh, trạng thái sống/mất); hiển thị thẻ FamilyClanCard chuyên nghiệp kèm bảng danh sách thành viên và ô tìm kiếm.")
    ]
    for title, desc in summary_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"• {title} ")
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r2 = p.add_run(desc)
        r2.font.size = Pt(9.5)

    # PHẦN 2: BẢNG TỔNG HỢP FILE
    add_h(doc, "2. BẢNG TỔNG HỢP TOÀN BỘ CÁC FILE ĐÃ THÊM MỚI VÀ SỬA ĐỔI", level=1)
    
    file_table_data = [
        ("STT", "Đường dẫn File", "Loại", "Mô tả trách nhiệm & Chức năng chính"),
        ("1", "BE/alembic/versions/0005_clan_join_requests.py", "Thêm mới", "Migration tạo bảng clan_join_requests và thêm cột lineage_type vào clan_memberships."),
        ("2", "BE/app/models/family/entities.py", "Sửa đổi", "Thêm model ClanJoinRequest và thuộc tính lineage_type vào ClanMembership."),
        ("3", "BE/app/models/family/__init__.py", "Sửa đổi", "Export ClanJoinRequest ra namespace của module family."),
        ("4", "BE/app/models/registry.py", "Sửa đổi", "Đăng ký ClanJoinRequest vào registry ORM SQLAlchemy toàn cục."),
        ("5", "BE/app/schemas/clan_join.py", "Thêm mới", "Khai báo Pydantic schemas cho API xin gia nhập, phê duyệt, danh sách dòng họ."),
        ("6", "BE/app/schemas/auth.py", "Sửa đổi", "Thêm danh sách joined_clans và active_clan_id vào schema phiên đăng nhập."),
        ("7", "BE/app/controllers/family_management/clan_join_router.py", "Thêm mới", "Router API REST cho gia nhập, rời dòng họ, danh sách dòng họ, phê duyệt hồ sơ."),
        ("8", "BE/app/controllers/family_management/clan_join_use_cases.py", "Thêm mới", "Xử lý nghiệp vụ logic kiểm tra, cấp membership, liên kết Person, rời dòng họ."),
        ("9", "BE/app/controllers/auth_access/use_cases.py", "Sửa đổi", "Bổ sung thông tin đa dòng họ kèm lineage_type khi cấp session và đăng nhập."),
        ("10", "BE/app/dependencies/permissions.py", "Sửa đổi", "Xác thực quyền hạn độc lập từng dòng họ; hỗ trợ vai trò BUSINESS_OWNER."),
        ("11", "BE/app/controllers/inter_family/router.py", "Sửa đổi", "Phân loại dòng họ nội/ngoại, truy vấn profile & danh sách thành viên Person thực tế."),
        ("12", "BE/app/main.py", "Sửa đổi", "Đăng ký router clan_join_router vào ứng dụng FastAPI chính."),
        ("13", "FE/src/types/auth.ts", "Sửa đổi", "Khai báo interface JoinedClan và kiểu quan hệ LineageType."),
        ("14", "FE/src/types/network.ts", "Sửa đổi", "Khai báo NetworkMemberItem, mở rộng FamilyNetwork chứa danh sách members."),
        ("15", "FE/src/types/tree.ts", "Sửa đổi", "Bổ sung cờ isUnlinked và unlinkedReason cho node trên cây gia phả."),
        ("16", "FE/src/services/clanJoin.service.ts", "Thêm mới", "Tầng client service gọi các API gia nhập, rời họ, danh sách dòng họ."),
        ("17", "FE/src/context/ClanContext.tsx", "Thêm mới", "React Context lưu trữ activeClanId, danh sách dòng họ, chuyển đổi cây đang xem."),
        ("18", "FE/src/components/dashboard/TopBar/ClanSelector.tsx", "Thêm mới", "Dropdown chọn dòng họ trên TopBar, hiển thị badge Nội/Ngoại, chống tràn text."),
        ("19", "FE/src/components/dashboard/TopBar/ClanSelector.css", "Thêm mới", "CSS chuyên biệt cho ClanSelector, hiệu ứng hover, badge màu sắc rõ ràng."),
        ("20", "FE/src/components/dashboard/TopBar/TopBar.tsx", "Sửa đổi", "Tích hợp ClanSelector, điều chỉnh layout chống đè chèn với ô tìm kiếm."),
        ("21", "FE/src/components/dashboard/TopBar/TopBar.css", "Sửa đổi", "Căn chỉnh gap 16px, flex responsive chống đè lấn giữa các khối trên TopBar."),
        ("22", "FE/src/components/profile/MyClans/MyClans.tsx", "Thêm mới", "Giao diện quản lý dòng tộc đã tham gia, nút rời họ, wizard 3 bước xin gia nhập."),
        ("23", "FE/src/components/profile/MyClans/MyClans.css", "Thêm mới", "CSS cho trang MyClans: grid dòng họ, bảng yêu cầu, modal wizard."),
        ("24", "FE/src/components/network/FamilyClanCard.tsx", "Thêm mới", "Thẻ hiển thị dòng họ chuyên nghiệp, bảng thành viên thực tế, avatar, chip đời."),
        ("25", "FE/src/components/network/FamilyPaternalTab/FamilyPaternalTab.tsx", "Sửa đổi", "Hiển thị danh sách thẻ FamilyClanCard cho Dòng họ Nội (Bên Cha)."),
        ("26", "FE/src/components/network/FamilyMaternalTab/FamilyMaternalTab.tsx", "Sửa đổi", "Hiển thị danh sách thẻ FamilyClanCard cho Dòng họ Ngoại (Bên Mẹ)."),
        ("27", "FE/src/components/network/Network.css", "Sửa đổi", "Bổ sung hệ thống CSS cho thẻ dòng họ, bảng thành viên, huy hiệu đời, nút toggle."),
        ("28", "FE/src/features/familyAdmin/FamilyAdminModule.tsx", "Sửa đổi", "Tích hợp tab duyệt yêu cầu gia nhập, thay window.confirm bằng modal bảo mật."),
        ("29", "FE/src/shared/tree/TreeLayout/TreeLayout.tsx", "Sửa đổi", "Xử lý hiển thị các node thành viên chưa nối phả hệ và nối tự động khi xác nhận."),
        ("30", "FE/src/shared/tree/TreeNodeCard/TreeNodeCard.tsx", "Sửa đổi", "Hiển thị giao diện đặc biệt cho node chưa nối phả hệ (viền cam đứt nét).")
    ]
    
    t_files = doc.add_table(rows=len(file_table_data), cols=4)
    widths = [Inches(0.5), Inches(2.6), Inches(0.9), Inches(2.8)]
    for r_idx, row in enumerate(file_table_data):
        for c_idx, val in enumerate(row):
            t_files.cell(r_idx, c_idx).text = val
    style_table(t_files, widths, header_bg="1E3A8A")
    
    # PHẦN 3: CHI TIẾT TỪNG FILE
    add_h(doc, "3. CHI TIẾT TỪNG FILE: MỤC ĐÍCH, NỘI DUNG SỬA ĐỔI & MÃ MINH HỌA", level=1)
    
    # 3.1 BACKEND
    add_h(doc, "3.1. Các File Phía Backend (Python / FastAPI / SQLAlchemy)", level=2)
    
    # File 1: router.py (inter_family)
    add_h(doc, "File: BE/app/controllers/inter_family/router.py (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Xử lý các endpoint mạng lưới liên họ, chia sẻ dữ liệu và liên kết giữa các dòng tộc (FR-BO-53..57).\n")
    p.add_run("Nội dung đã sửa đổi:\n").bold = True
    p.add_run("1. Trước đây: Endpoint ").font.size = Pt(9.5)
    p.add_run("GET /networks/families").bold = True
    p.add_run(" chỉ trả về danh sách cứng (mock) toàn bộ dòng tộc có trong DB với category='thong_gia', không kiểm tra tham số category ('noi'/'ngoai'), không có thông tin thủy tổ, lịch sử và không có thành viên.\n")
    p.add_run("2. Đã sửa: Viết lại hoàn chỉnh hàm ").font.size = Pt(9.5)
    p.add_run("get_network_families").bold = True
    p.add_run(":\n"
              "   - Phân tích memberships của user hiện tại (account.user_id) để tách paternal_clan_ids (bên nội) và maternal_clan_ids (bên ngoại).\n"
              "   - Kiểm tra liên kết InterClanLink đang ACTIVE giữa các dòng họ.\n"
              "   - Lọc chính xác category: 'noi' chỉ lấy dòng họ nội của user, 'ngoai' lấy dòng họ ngoại và dòng họ liên kết thông gia/hôn nhân.\n"
              "   - Truy vấn ClanProfile để lấy origin_place, history, founder_person_id.\n"
              "   - Xác định founderName (qua founder_person_id hoặc người thuộc Đời 1 / generation=1 ưu tiên nam giới).\n"
              "   - Truy vấn toàn bộ Person thuộc dòng tộc, format mảng members gồm: id, fullName, gender (Nam/Nữ/Khác), generation, birthDate, deathDate, isAlive, branch, bio.")
    
    add_code_block(doc, 
"""# BE/app/controllers/inter_family/router.py - get_network_families (Đã sửa)
@router.get("/families")
async def get_network_families(
    category: str = Query("all"),
    db: AsyncSession = Depends(get_db),
    account: Account = Depends(get_current_account),
):
    # 1. Xác định quan hệ của user
    user_memberships = (await db.execute(
        select(ClanMembership).where(ClanMembership.user_id == account.user_id, ClanMembership.status == "ACTIVE")
    )).scalars().all()
    user_clan_ids = {m.clan_id for m in user_memberships}
    maternal_clan_ids = {m.clan_id for m in user_memberships if (m.lineage_type or "").upper() == "MATERNAL"}
    paternal_clan_ids = user_clan_ids - maternal_clan_ids

    # 2. Lọc theo category ('noi' vs 'ngoai')
    # 3. Batch query ClanProfile, Branch, và Person
    # 4. Trả về đầy đủ founderName, originPlace, history, memberCount, và members list
""")

    # File 2: entities.py (family)
    add_h(doc, "File: BE/app/models/family/entities.py (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Định nghĩa các thực thể ORM cho Dòng họ (Clan), Thành viên (ClanMembership), Hồ sơ (ClanProfile).\n")
    p.add_run("Nội dung đã sửa đổi:\n").bold = True
    p.add_run("1. Thêm cột ").font.size = Pt(9.5)
    p.add_run("lineage_type = mapped_column(String(30), nullable=True)").bold = True
    p.add_run(" trong model ClanMembership để phân biệt người dùng tham gia nhánh Nội ('PATERNAL') hay Ngoại ('MATERNAL').\n")
    p.add_run("2. Thêm thực thể ORM mới ").font.size = Pt(9.5)
    p.add_run("ClanJoinRequest").bold = True
    p.add_run(" ánh xạ bảng clan_join_requests, lưu trữ thông tin: requesting_user_id, target_clan_id, target_lineage_type, relative_name, relationship_desc, status ('PENDING', 'APPROVED', 'REJECTED', 'NEED_INFO'), reviewed_by_user_id, review_notes, reviewed_at.")

    # File 3: permissions.py
    add_h(doc, "File: BE/app/dependencies/permissions.py (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Phân quyền và bảo vệ các API endpoint dựa trên vai trò người dùng trong dòng họ.\n")
    p.add_run("Nội dung đã sửa đổi:\n").bold = True
    p.add_run("Cập nhật logic hàm ").font.size = Pt(9.5)
    p.add_run("is_active_owner").bold = True
    p.add_run(" và kiểm tra quyền admin: kiểm tra theo từng clan_id cụ thể thay vì cấp quyền toàn cục; đồng thời công nhận người dùng sở hữu vai trò BUSINESS_OWNER là quản trị viên hợp lệ của dòng họ.")

    # File 4: clan_join_router.py & use_cases.py
    add_h(doc, "File: BE/app/controllers/family_management/clan_join_router.py & use_cases.py (THÊM MỚI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Cung cấp toàn bộ API RESTful và nghiệp vụ quản lý gia nhập dòng họ và rời dòng họ.\n")
    p.add_run("Các endpoint triển khai:\n").bold = True
    p.add_run("• POST /clans/join-requests: Gửi yêu cầu gia nhập kèm khai báo quan hệ.\n"
              "• GET /clans/my-requests: Lấy danh sách yêu cầu bản thân đã gửi kèm phản hồi người duyệt.\n"
              "• GET /clans/my-clans: Lấy các dòng tộc đã tham gia, trạng thái, vai trò và phân loại nội/ngoại.\n"
              "• GET /clans/available: Tìm kiếm dòng tộc công khai để xin gia nhập.\n"
              "• POST /clans/leave: Cho phép người dùng rời dòng họ nếu vào nhầm.\n"
              "• GET /clans/join-requests/pending: Dành cho Family Admin xem hồ sơ chờ duyệt.\n"
              "• POST /clans/join-requests/{id}/review: Duyệt hoặc từ chối hồ sơ, tự động tạo ClanMembership với quyền MEMBER và liên kết Person.")

    # 3.2 FRONTEND
    add_h(doc, "3.2. Các File Phía Frontend (React / TypeScript / CSS)", level=2)
    
    # File 5: MyClans.tsx
    add_h(doc, "File: FE/src/components/profile/MyClans/MyClans.tsx (THÊM MỚI & SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Trang quản lý trung tâm 'Dòng tộc của tôi' cho phép người dùng quản lý các cây đã tham gia, đổi cây, gửi yêu cầu xin gia nhập dòng họ mới, rời dòng họ.\n")
    p.add_run("Nội dung đã sửa đổi (Dọn dẹp icon & Tối ưu UI):\n").bold = True
    p.add_run("• Dòng 335: Xoá icon ").font.size = Pt(9.5)
    p.add_run("🏛️").bold = True
    p.add_run(" khỏi tiêu đề chính -> ").font.size = Pt(9.5)
    p.add_run("<h1>Dòng tộc của tôi</h1>\n").bold = True
    p.add_run("• Dòng 341: Xoá icon ").font.size = Pt(9.5)
    p.add_run("🔄").bold = True
    p.add_run(" khỏi nút làm mới -> ").font.size = Pt(9.5)
    p.add_run("Làm mới\n").bold = True
    p.add_run("• Dòng 353: Xoá icon ").font.size = Pt(9.5)
    p.add_run("🌳").bold = True
    p.add_run(" khỏi tiêu đề mục dòng tộc đã tham gia -> ").font.size = Pt(9.5)
    p.add_run("<h2>Dòng tộc đã tham gia</h2>\n").bold = True
    p.add_run("• Dòng 436: Xoá icon ").font.size = Pt(9.5)
    p.add_run("📩").bold = True
    p.add_run(" khỏi tiêu đề bảng yêu cầu -> ").font.size = Pt(9.5)
    p.add_run("<h2>Yêu cầu gia nhập dòng tộc</h2>\n").bold = True
    p.add_run("• Dòng 773: Xoá icon ").font.size = Pt(9.5)
    p.add_run("ℹ️").bold = True
    p.add_run(" khỏi khối lưu ý quy trình xác minh.\n")
    p.add_run("• Tích hợp modal Rời dòng họ an toàn: khi click 'Rời dòng họ', hiển thị cảnh báo xác nhận trước khi thực hiện API POST /clans/leave.")

    # File 6: FamilyClanCard.tsx
    add_h(doc, "File: FE/src/components/network/FamilyClanCard.tsx (THÊM MỚI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích của file: ").bold = True
    p.add_run("Component chuyên biệt phục vụ hiển thị thẻ thông tin dòng tộc và danh sách thành viên thực tế trong mục Dòng họ Nội và Dòng họ Ngoại.\n")
    p.add_run("Cấu trúc và tính năng của component:\n").bold = True
    p.add_run("1. Khối Header Dòng Tộc: Hiển thị tên dòng họ, mã dòng họ (CLAN-VO-01, CLAN-MAI-02), tag phân loại (Bên nội / Bên ngoại), tag trạng thái kết nối.\n"
              "2. Lưới thông tin cốt lõi: Thủy tổ (in đậm nổi bật), Quê quán, Từ đường, Lịch sử tóm tắt dòng họ.\n"
              "3. Bảng Thành Viên Dòng Tộc (Members Table):\n"
              "   - Cột Họ và tên: Avatar chữ cái phân biệt giới tính (xanh cho Nam, hồng cho Nữ), tên thành viên, chi nhánh trực thuộc.\n"
              "   - Cột Đời: Huy hiệu màu sắc theo đời (Đời 1, Đời 2, Đời 3, Đời 4).\n"
              "   - Cột Giới tính: Nam / Nữ.\n"
              "   - Cột Ngày sinh: Định dạng DD/MM/YYYY chuẩn mực.\n"
              "   - Cột Trạng thái: Huy hiệu 'Còn sống' (xanh lá) hoặc 'Đã mất' (xám).\n"
              "4. Tính năng tìm kiếm thành viên: Ô input 'Tìm thành viên, đời...' lọc trực tiếp theo thời gian thực.\n"
              "5. Nút 'Xem tất cả X thành viên' / 'Thu gọn': Thu gọn danh sách nếu dòng họ có nhiều thành viên (mặc định hiện 5 người đầu tiên).")

    # File 7 & 8: FamilyPaternalTab & FamilyMaternalTab
    add_h(doc, "File: FamilyPaternalTab.tsx & FamilyMaternalTab.tsx (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích & Thay đổi:\n").bold = True
    p.add_run("• Thay thế đoạn mã render các card sơ sài cũ bằng component ").font.size = Pt(9.5)
    p.add_run("FamilyClanCard").bold = True
    p.add_run(".\n• FamilyPaternalTab gọi ").font.size = Pt(9.5)
    p.add_run("fetchFamilyNetwork('noi')").bold = True
    p.add_run(" -> Hiển thị Võ Đại Tộc (Gò Nổi - Quảng Nam) với đầy đủ 15 thành viên.\n"
              "• FamilyMaternalTab gọi ").font.size = Pt(9.5)
    p.add_run("fetchFamilyNetwork('ngoai')").bold = True
    p.add_run(" -> Hiển thị Mai Thế Tộc (Phong Điền - Thừa Thiên Huế) với đầy đủ 19 thành viên.")

    # File 9: Network.css
    add_h(doc, "File: FE/src/components/network/Network.css (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích & Thay đổi:\n").bold = True
    p.add_run("Bổ sung hơn 200 dòng CSS hiện đại cho .network-clan-card, .network-clan-header-top, .network-clan-info-grid, .network-members-table-wrap, .member-avatar-badge, .gen-pill (gen-1 đến gen-4), .living-badge, .btn-toggle-members, đảm bảo trải nghiệm giao diện cao cấp, chuyên nghiệp.")

    # File 10: TopBar & ClanSelector
    add_h(doc, "File: ClanSelector.tsx & TopBar.tsx / TopBar.css (THÊM MỚI & SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích & Thay đổi:\n").bold = True
    p.add_run("• Xây dựng bộ chọn dòng họ ClanSelector trên thanh điều hướng đầu trang (TopBar).\n"
              "• Khắc phục triệt để lỗi đè chèn giữa bộ chọn dòng họ và thanh tìm kiếm: thêm ").font.size = Pt(9.5)
    p.add_run("flex-shrink: 0, text-overflow: ellipsis, max-width: 220px").bold = True
    p.add_run(" cho ClanSelector; thêm ").font.size = Pt(9.5)
    p.add_run("gap: 16px và flex: 1 1 auto").bold = True
    p.add_run(" cho thanh tìm kiếm.")

    # File 11: TreeLayout & TreeNodeCard
    add_h(doc, "File: TreeLayout.tsx & TreeNodeCard.tsx (SỬA ĐỔI)", level=3)
    p = doc.add_paragraph()
    p.add_run("Mục đích & Thay đổi:\n").bold = True
    p.add_run("• Bổ sung logic hiển thị tài khoản đã tham gia dòng họ nhưng chưa xác nhận quan hệ với ai: thành viên này được render trên sơ đồ cây gia phả với trạng thái đặc biệt (unlinked).\n"
              "• Giao diện TreeNodeCard hiển thị viền đứt nét màu cam hổ phách, badge cảnh báo 'Chờ nối quan hệ' và không vẽ đường liên kết phả hệ cho đến khi có xác nhận quan hệ chính thức.")

    # PHẦN 4: BẰNG CHỨNG KIỂM THỬ THỰC TẾ
    add_h(doc, "4. BẰNG CHỨNG KIỂM THỬ THỰC TẾ & KẾT QUẢ ĐẠT ĐƯỢC", level=1)
    
    add_callout(
        doc,
        "KẾT QUẢ KIỂM THỬ TÍCH HỢP TỰ ĐỘNG (API VERIFICATION)",
        "Kiểm thử tự động gọi thực tế vào API backend sau khi cập nhật:\n\n"
        "1. DÒNG HỌ NỘI (category=noi):\n"
        "   - HTTP Status: 200 OK\n"
        "   - Dòng tộc trả về: Võ Đại Tộc (Gò Nổi - Quảng Nam) [Mã: CLAN-VO-01]\n"
        "   - Thủy tổ: Cụ Võ Văn Tiên | Quê quán: Gò Nổi, Huyện Điện Bàn, Tỉnh Quảng Nam\n"
        "   - Số lượng thành viên thực tế: 15 người\n"
        "   - Mẫu thành viên: Cụ Võ Văn Tiên (Đời 1), Cụ Bà Lê Thị Lành (Đời 1), Võ Văn Trung (Đời 2), Trần Thị Huệ (Đời 2), Võ Trí Thắng (Đời 3), Võ Minh Khôi (Đời 4)...\n\n"
        "2. DÒNG HỌ NGOẠI (category=ngoai):\n"
        "   - HTTP Status: 200 OK\n"
        "   - Dòng tộc trả về: Mai Thế Tộc (Phong Điền - Thừa Thiên Huế) [Mã: CLAN-MAI-02]\n"
        "   - Thủy tổ: Cụ Mai Thế Hùng | Quê quán: Phong Điền, Tỉnh Thừa Thiên Huế\n"
        "   - Số lượng thành viên thực tế: 19 người\n"
        "   - Mẫu thành viên: Cụ Mai Thế Hùng (Đời 1), Cụ Bà Hoàng Thị Nhàn (Đời 1), Mai Thế Đức (Đời 2), Mai Phúc Hậu (Đời 3), Trần Minh Quân (Đời 4), Võ Bích Hạnh (Đời 4)...",
        bg_hex="F8FAFC",
        border_hex="0284C7"
    )

    # Save document
    output_path = "/Users/vanthang/Documents/NCKH-CapTone1/Gia-ph-/Bao_cao_chi_tiet_cac_file_them_sua_va_tinh_nang.docx"
    doc.save(output_path)
    print(f"File created successfully at: {output_path}")

if __name__ == "__main__":
    main()
